"""SIH 2026 idea deck for eRTMAC NWIS, following the official template layout."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt, Emu

D = '/tmp/claude-0/-home-user-eRTMAC-NWIS/6574508b-7bfc-5d92-9894-fa2cea316fe8/scratchpad/'
SIH_LOGO = D + 'tpl_2_0_X13.png'
APP_LOGO = D + 'ppt/logo.png'
TEAM = 'COOLKATS'
W, H = Pt(1440), Pt(810)

NAVY = RGBColor(0x1F, 0x49, 0x7D)
BLUE = RGBColor(0x00, 0x70, 0xC0)
LINK = RGBColor(0x1F, 0x4E, 0x8C)
INK = RGBColor(0x1E, 0x1E, 0x1E)
GREY = RGBColor(0x59, 0x59, 0x59)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
PURPLE = RGBColor(0x80, 0x64, 0xA2)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
GREEN = RGBColor(0x13, 0x8A, 0x36)
BROWN = RGBColor(0x2B, 0x26, 0x23)
SAND = RGBColor(0xF6, 0xEF, 0xE2)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
blank = prs.slide_layouts[6]


def text(slide, x, y, w, h, runs, size=22, font='Arial', color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, {opts})."""
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, Pt(2))
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get('font', font)
            f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold)
            f.underline = o.get('underline', False)
            f.color.rgb = o.get('color', color)
        if isinstance(para, list) and para and para[0][1].get('space'):
            p.space_before = Pt(para[0][1]['space'])
    return tb


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.5, radius=0.12):
    s = slide.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s


def arrow(slide, x1, y1, x2, y2, color=GREY):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
    c.line.color.rgb = color
    c.line.width = Pt(2.25)
    ln = c.line._get_or_add_ln()
    from lxml import etree
    tail = etree.SubElement(ln, '{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd')
    tail.set('type', 'triangle')
    return c


def picture(slide, path, x, y, w=None, h=None, border=True):
    pic = slide.shapes.add_picture(path, Pt(x), Pt(y), Pt(w) if w else None, Pt(h) if h else None)
    if border:
        pic.line.color.rgb = RGBColor(0xD0, 0xD0, 0xD0)
        pic.line.width = Pt(1)
    return pic


def frame(n, title):
    """Template chrome: team-name oval, centred title, SIH logo, blue footer bar."""
    s = prs.slides.add_slide(blank)
    ov = box(s, 34, 26, 150, 78, line=PURPLE, shape=MSO_SHAPE.OVAL, lw=2)
    text(s, 34, 26, 150, 78, [[(TEAM, {'bold': True, 'size': 20})]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, font='Calibri')
    text(s, 230, 52, 980, 70, [[(title, {'bold': True})]], size=44, font='Times New Roman', align=PP_ALIGN.CENTER)
    s.shapes.add_picture(SIH_LOGO, Pt(1150), Pt(8), Pt(240))
    bar = box(s, 0, 750, 1440, 60, fill=BLUE, shape=MSO_SHAPE.RECTANGLE)
    text(s, 420, 762, 600, 36, '@SIH Idea submission- Template', size=16, color=RGBColor(255, 255, 255), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1320, 762, 60, 36, [[(str(n), {'bold': True})]], size=16, color=RGBColor(255, 255, 255), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return s


def heading(s, x, y, w, t, size=24):
    text(s, x, y, w, 40, [[('• ', {'color': LINK, 'bold': True}), (t, {'bold': True, 'underline': True, 'color': LINK})]], size=size)


def bullets(s, x, y, w, h, items, size=17, gap=5):
    paras = []
    for it in items:
        if isinstance(it, tuple):
            head, rest = it
            paras.append([('• ', {'bold': True, 'space': gap}), (head, {'bold': True}), (rest, {})])
        else:
            paras.append([('• ', {'bold': True, 'space': gap}), (it, {})])
    return text(s, x, y, w, h, paras, size=size)


# ---------------------------------------------------------------- 1. Title
s = prs.slides.add_slide(blank)
text(s, 150, 18, 1000, 80, [[('SMART INDIA HACKATHON 2026', {'bold': True})]], size=56, font='Garamond', color=NAVY, align=PP_ALIGN.CENTER)
s.shapes.add_picture(SIH_LOGO, Pt(1150), Pt(8), Pt(240))
for (x, y, w, h) in [(960, 175, 520, 470), (1020, 105, 170, 150), (825, 445, 150, 130)]:
    hx = box(s, x, y, w, h, fill=LIGHT, shape=MSO_SHAPE.HEXAGON)
text(s, 300, 118, 600, 70, [[('TITLE PAGE', {'bold': True})]], size=46, font='Times New Roman', align=PP_ALIGN.CENTER)
rows = [('Problem Statement ID – ', '26121'),
        ('Problem Statement Title- ', 'eRTMAC NWIS: AI-powered offset well knowledge and decision support platform for drilling operations'),
        ('Theme- ', 'Smart Automation'),
        ('PS Category- ', 'Software'),
        ('Team ID- ', '187666'),
        ('Team Name- ', 'COOLKATS')]
y = 225
for head, val in rows:
    h = 92 if 'Title' in head else 62
    text(s, 70, y, 900, h, [[('• ', {'bold': True}), (head, {'bold': True}), (val, {'bold': False, 'color': NAVY})]], size=26)
    y += h + 8
bulb = s.shapes.add_picture(D + 'tpl_1_1_X12.png', Pt(1010), Pt(200), Pt(342), Pt(385))
bulb.crop_right = 0.58

# ---------------------------------------------------------------- 2. Idea
s = frame(2, 'IDEA TITLE: eRTMAC NWIS')
heading(s, 40, 120, 1100, 'Proposed Solution (Describe your Idea/Solution/Prototype)')
text(s, 60, 160, 820, 56, [[('One map-first platform that turns old offset-well reports into ', {}), ('early warnings for the well being drilled today.', {'bold': True})]], size=18)
bullets(s, 60, 212, 800, 300, [
    ('Location intelligence map: ', 'tap any spot → history, rock & soil, success score, hazards (landslide, flood, subsidence, gas kick), legal / illegal zone, estimated oil (P90–P10) and land ownership.'),
    ('Offset well intelligence: ', 'wells within a chosen radius, searchable events (mud loss, kick, stuck pipe, torque, cementing, NPT) with cause, fix and lesson; depth-vs-formation view.'),
    ('Alert engine: ', 'matches offset problems to the active well by formation and warns ~150 m before the bit gets there, with “what worked nearby”.'),
    ('Driller registration + live geotag: ', 'documents approved by admin, rigs declared; allowed zone = rig reach + safety margin; breach of zone or forest/protected area is flagged instantly.'),
    ('AI document extraction: ', 'PDF drilling reports → OCR → LLM/rules → structured events, lessons and formation tops.'),
], size=16, gap=6)
heading(s, 40, 560, 600, 'How it addresses the problem', size=20)
bullets(s, 60, 594, 800, 190, [
    'Knowledge buried in DDRs, WCRs and engineers’ memory becomes searchable in seconds.',
    'Correlates drilling, geology and reservoir data by depth and formation across wells.',
    'Field view on phone/tablet for rig crews, office view with analytics and approvals.',
], size=16)
picture(s, D + 'ppt/s_map.png', 885, 130, 520)
picture(s, D + 'ppt/s_field.png', 1262, 440, h=300)
heading(s, 885, 440, 380, 'Innovation and uniqueness', size=19)
bullets(s, 900, 476, 350, 270, [
    'Formation-based depth correlation, not raw depth',
    'Untapped-spot layer: ML score + confidence',
    'Legality + land ownership in the same tap',
    'Physics for range: haversine + minimum-curvature bottom-hole',
    'Works offline with a built-in India basemap',
], size=14, gap=3)

# ---------------------------------------------------------------- 3. Technical approach
s = frame(3, 'TECHNICAL APPROACH')
heading(s, 40, 120, 800, 'Technologies to be used', size=22)
tech = [('Frontend', 'React + Vite, Leaflet maps, Recharts, light/dark theme, PWA-style field view'),
        ('Backend', 'Python FastAPI, WebSockets (live tracking & alerts), JWT roles: driller / engineer / admin'),
        ('Data', 'SQLite → PostgreSQL + PostGIS in production; GEM tracker, Volve DDR, GSI Bhukosh, SoilGrids, WDPA, NASA GLC'),
        ('AI / ML', 'scikit-learn random forests (success score, risk per problem type); OCR (Tesseract) + Claude LLM with rule-based fallback'),
        ('Physics', 'Haversine, point-in-polygon, minimum-curvature trajectory, volumetric STOIIP with Monte Carlo')]
y = 162
for k, v in tech:
    box(s, 60, y, 150, 50, fill=BROWN)
    text(s, 60, y, 150, 50, [[(k, {'bold': True})]], size=16, color=SAND, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 222, y, 560, 54, v, size=14.5, anchor=MSO_ANCHOR.MIDDLE)
    y += 60
heading(s, 40, 470, 800, 'Methodology and process for implementation', size=22)
steps = ['Ingest\nPDF / DDR / GIS', 'Extract\nOCR + LLM', 'Store\nwells, events,\nzones', 'Analyse\nML + formation\ncorrelation', 'Act\nmap, alerts,\ntracking']
x = 60
for i, st in enumerate(steps):
    c = [NAVY, BLUE, GREEN, ORANGE, BROWN][i]
    box(s, x, 515, 140, 110, fill=c)
    t, rest = st.split('\n', 1)
    text(s, x + 4, 520, 132, 100, [[(t, {'bold': True, 'size': 17})], [(rest.replace('\n', ' '), {'size': 12.5})]], color=RGBColor(255, 255, 255), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < 4:
        arrow(s, x + 142, 570, x + 160, 570)
    x += 162
text(s, 60, 640, 780, 100, [[('Live loop: ', {'bold': True}), ('bit depth updates → alert engine re-runs → WebSocket pushes warning + recommendation to field & office. Driller GPS → range check → breach logged with time & coordinates.', {})]], size=14.5)
picture(s, D + 'ppt/s_offset.png', 900, 128, h=275)
picture(s, D + 'ppt/s_tracking.png', 900, 438, h=275)
text(s, 900, 406, 470, 24, 'Offset-well correlation view (depth vs formation, problems as dots)', size=12, color=GREY)
text(s, 900, 716, 470, 20, 'Live tracking: allowed zone, breach log', size=12, color=GREY)

# ---------------------------------------------------------------- 4. Feasibility
s = frame(4, 'FEASIBILITY AND VIABILITY')
cols = [
    ('Analysis of the feasibility of the idea', GREEN, [
        'Working prototype already built and deployed (web + phone view)',
        'Open-source stack – no licence cost; runs on one server or Docker',
        'Plugs into eRTMAC as a knowledge layer; public datasets exist for all layers',
        'Loaders ready for real data (GEM, Volve, Bhukosh, WDPA, SRTM)',
        '36 automated tests on geometry, alerts, APIs, extraction']),
    ('Potential challenges and risks', ORANGE, [
        'Old reports are scanned, handwritten or inconsistent',
        'Real drilling data is confidential; limited training data',
        'Poor connectivity at remote rig sites',
        'False alerts reduce trust of drilling engineers',
        'Land-record and zone boundaries are not digital everywhere']),
    ('Strategies for overcoming these challenges', BLUE, [
        'OCR + LLM with human review; rule-based fallback always available',
        'On-premise deployment inside OIL network, role-based access',
        'Offline-capable field view, built-in basemap, sync when online',
        'Every alert shows its evidence wells; engineers confirm or dismiss',
        'Start with DGH NDR + state land portals; flag “needs licence” when unknown']),
]
x = 45
for title, c, items in cols:
    box(s, x, 130, 440, 58, fill=c)
    text(s, x, 130, 440, 58, [[(title, {'bold': True})]], size=18, color=RGBColor(255, 255, 255), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    box(s, x, 194, 440, 540, fill=RGBColor(0xFA, 0xFA, 0xFA), line=RGBColor(0xDD, 0xDD, 0xDD), lw=1)
    bullets(s, x + 16, 210, 410, 520, items, size=20, gap=20)
    x += 455

# ---------------------------------------------------------------- 5. Impact
s = frame(5, 'IMPACT AND BENEFITS')
heading(s, 40, 120, 900, 'Potential impact on the target audience', size=22)
aud = [('Drilling engineers', 'warnings before trouble zones; lessons from nearby wells in one tap'),
       ('Rig crews (field)', 'simple phone view: current well, alerts, live location, risk ahead'),
       ('Office & management', 'all wells, drillers, NPT trends, approvals in one dashboard'),
       ('Regulators / admin', 'verified drillers, documented rigs, zone breaches logged')]
y = 162
for k, v in aud:
    text(s, 60, y, 820, 40, [[('▸ ', {'color': ORANGE, 'bold': True}), (k + ': ', {'bold': True}), (v, {})]], size=17)
    y += 44
heading(s, 40, 355, 900, 'Benefits of the solution (social, economic, environmental)', size=22)
ben = [('Economic', GREEN, 'Less NPT from losses, kicks and stuck pipe; faster decisions; better well placement with untapped-spot scoring'),
       ('Safety', ORANGE, 'Kick / overpressure alerts ahead of the bit; earthquake, flood and landslide checks at site selection'),
       ('Environmental', BLUE, 'No drilling inside forests, wetlands or protected areas (e.g. Baghjan 2020 lesson); breach alerts'),
       ('Social / knowledge', NAVY, 'Expert knowledge preserved when engineers retire; land ownership shown before work starts')]
x, y = 60, 398
for i, (k, c, v) in enumerate(ben):
    xx = 60 + (i % 2) * 420
    yy = 398 + (i // 2) * 170
    box(s, xx, yy, 400, 155, fill=RGBColor(0xFA, 0xFA, 0xFA), line=c, lw=2)
    text(s, xx + 16, yy + 10, 370, 140, [[(k, {'bold': True, 'size': 19, 'color': c})], [(v, {'size': 15, 'space': 4})]])
picture(s, D + 'ppt/s_overview.png', 910, 130, 495)
text(s, 910, 430, 495, 22, 'Office dashboard: KPIs, alerts, NPT by formation, trends', size=12, color=GREY)
box(s, 910, 470, 495, 264, fill=SAND)
text(s, 930, 484, 460, 240, [[('What changes for OIL', {'bold': True, 'size': 20})],
     [('Before: ', {'bold': True, 'space': 10}), ('hours searching PDFs and asking seniors', {})],
     [('After: ', {'bold': True, 'space': 6}), ('offset lessons and risks in seconds, pushed live to the rig', {})],
     [('Before: ', {'bold': True, 'space': 10}), ('site legality & ownership checked manually', {})],
     [('After: ', {'bold': True, 'space': 6}), ('one tap on the map', {})]], size=16)

# ---------------------------------------------------------------- 6. References
s = frame(6, 'RESEARCH AND REFERENCES')
refs = [
    ('Problem & domain', ['SIH 2026 PS 26121 – Oil India Limited, eRTMAC NWIS',
                          'DGH National Data Repository – ndr.dghindia.gov.in',
                          'Mitchell & Miska, Fundamentals of Drilling Engineering (SPE) – minimum curvature, well control']),
    ('Datasets', ['Global Energy Monitor – Global Oil & Gas Extraction Tracker',
                  'Equinor Volve field data (daily drilling reports) – equinor.com/energy/volve-data-sharing',
                  'GSI Bhukosh geology – bhukosh.gsi.gov.in',
                  'ISRIC SoilGrids v2.0 – rest.isric.org',
                  'World Database on Protected Areas – protectedplanet.net',
                  'NASA Global Landslide Catalog; OpenTopography SRTM 30 m',
                  'Natural Earth (India point-of-view boundaries); OpenStreetMap']),
    ('Tools', ['FastAPI, React, Leaflet, scikit-learn, Tesseract OCR, Anthropic Claude API']),
]
y = 128
heading(s, 40, 118, 860, 'Details / Links of the reference and research work', size=22)
y = 165
for head, items in refs:
    text(s, 60, y, 800, 34, [[(head, {'bold': True})]], size=19)
    y += 38
    bullets(s, 60, y, 830, 40 * len(items), items, size=15, gap=2)
    y += 24 * len(items) + 30
box(s, 930, 130, 470, 600, fill=SAND)
text(s, 950, 150, 430, 560, [[('Prototype', {'bold': True, 'size': 22})],
     [('Live demo: ', {'bold': True, 'space': 10}), ('gurltff.github.io/eRTMAC-NWIS', {'color': LINK})],
     [('Code: ', {'bold': True, 'space': 6}), ('github.com/gurltff/eRTMAC-NWIS', {'color': LINK})],
     [('Demo logins: ', {'bold': True, 'space': 14}), ('engineer@nwis.demo / Engineer@123', {})],
     [('driller@nwis.demo / Driller@123', {'space': 2})],
     [('Note: ', {'bold': True, 'space': 14}), ('prototype runs on sample data for Upper Assam; field and formation names are real, all values illustrative.', {'color': GREY})]], size=15)

out = D + 'ppt/COOLKATS_187666_SIH2026_eRTMAC-NWIS.pptx'
prs.save(out)
print(out)
